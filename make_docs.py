#!/usr/bin/env python3
"""
Build the Operating Review deck (PPTX) and document (DOCX) from index.html.

Both are generated from ONE parse of the page, so deck, document and site cannot drift.
Styling reproduces the site's own theme, read from its :root custom properties:
    --bg  #FAF9F5   warm ground        --tx  #1F1E1D   ink
    --alt #F4F2EC   panel              --tx2 #5C5A54   muted
    --ln  #E4E0D5   rule               --tx3 #8A8781   light
    --ac  #D97757   terracotta accent
Logo assets come from make_logo.py, which reproduces the .mark CSS exactly.
"""
import re, html, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
import docx
from docx.shared import Pt as DPt, RGBColor as DRGB, Inches as DIn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE   = "/Users/lennyshovsky/Downloads/wirewalkai/"
SRC    = BASE + "index.html"
PPTX   = BASE + "operating-review.pptx"
DOCX   = BASE + "operating-review.docx"
LOCKUP = BASE + "logo-lockup.png"
MARK   = BASE + "logo-mark.png"

BG   = RGBColor(0xFA, 0xF9, 0xF5)
PANEL= RGBColor(0xF4, 0xF2, 0xEC)
RULE = RGBColor(0xE4, 0xE0, 0xD5)
TX   = RGBColor(0x1F, 0x1E, 0x1D)
TX2  = RGBColor(0x5C, 0x5A, 0x54)
TX3  = RGBColor(0x8A, 0x87, 0x81)
AC   = RGBColor(0xD9, 0x77, 0x57)
D_TX, D_TX2, D_TX3, D_AC = DRGB(0x1F,0x1E,0x1D), DRGB(0x5C,0x5A,0x54), DRGB(0x8A,0x87,0x81), DRGB(0xD9,0x77,0x57)


def parse(path):
    s = open(path, encoding="utf-8").read()
    s = re.sub(r"<(script|style|noscript)[^>]*>.*?</\1>", "", s, flags=re.S | re.I)

    # Some blocks on the page carry their meaning in <b>/<span> rather than in a
    # heading and a paragraph, so a tag-based reader walks straight past them.
    # Two of those are load-bearing: the RECOVERABLE / REPEATING / CONTINGENT
    # horizons under each loss, and the scoping offer under AI adoption. Both
    # were silently absent from the deck and the document until the output was
    # checked against the page rather than against the file size.
    #
    # Rewritten into <p> before parsing rather than by adding <b> and <span> to
    # the tag list, which would drag in every inline emphasis on the page.
    s = re.sub(r'<div class="gets"><b>(.*?)</b>\s*<span>(.*?)</span></div>',
               r"<p>\1 — \2</p>", s, flags=re.S)
    s = re.sub(r'<div class="aibuy">\s*<div>\s*<b>(.*?)</b>\s*<span>(.*?)</span>',
               r"<p>\1 — \2</p>", s, flags=re.S)
    out = []
    for m in re.finditer(r'<section[^>]*id="([^"]+)"[^>]*>(.*?)</section>', s, re.S | re.I):
        sid, body = m.group(1), m.group(2)
        blocks = []
        for b in re.finditer(r"<(h1|h2|h3|h4|p|li)[^>]*>(.*?)</\1>", body, re.S | re.I):
            tag = b.group(1).lower()
            txt = re.sub(r"<[^>]+>", " ", b.group(2))
            txt = html.unescape(re.sub(r"\s+", " ", txt)).strip()
            if txt and len(txt) > 1:
                blocks.append((tag, txt))
        if blocks:
            out.append((sid, blocks))
    return out


def heading_of(blocks):
    for t, x in blocks:
        if t in ("h1", "h2"):
            return x
    return ""


def grouped(blocks, lead):
    groups, cur = [], None
    for t, x in blocks:
        if t in ("h3", "h4"):
            cur = (x, []); groups.append(cur)
        elif t == "li":
            if cur is None:
                cur = ("", []); groups.append(cur)
            cur[1].append(x)
        elif t == "p" and cur is not None and x != lead:
            cur[1].append(x)
    return groups


# ------------------------------------------------------------------ PPTX
def build_pptx(sections):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def new_slide():
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = BG
        return sl

    def tb(sl, l, t, w, h):
        box = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        box.text_frame.word_wrap = True
        return box.text_frame

    def rule(sl, l, t, w):
        ln = sl.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Pt(1.4))
        ln.fill.solid(); ln.fill.fore_color.rgb = AC; ln.line.fill.background()
        ln.shadow.inherit = False

    # ---- title slide
    sid, blocks = sections[0]
    sl = new_slide()
    if os.path.exists(LOCKUP):
        sl.shapes.add_picture(LOCKUP, Inches(0.85), Inches(0.75), height=Inches(0.46))
    rule(sl, 0.9, 2.15, 1.5)
    tf = tb(sl, 0.85, 2.45, 11.4, 2.3)
    r = tf.paragraphs[0].add_run(); r.text = heading_of(blocks)
    r.font.size, r.font.bold, r.font.color.rgb = Pt(38), True, TX
    lead = next((x for t, x in blocks if t == "p"), "")
    if lead:
        p = tf.add_paragraph(); p.space_before = Pt(20)
        r2 = p.add_run(); r2.text = lead
        r2.font.size, r2.font.color.rgb = Pt(14.5), TX2
    tf3 = tb(sl, 0.85, 6.45, 11.4, 0.5)
    r3 = tf3.paragraphs[0].add_run(); r3.text = "The Operating Review"
    r3.font.size, r3.font.color.rgb, r3.font.bold = Pt(11.5), AC, True

    # ---- content slides
    for sid, blocks in sections[1:]:
        head = heading_of(blocks)
        lead = next((x for t, x in blocks if t == "p"), "")
        groups = grouped(blocks, lead)

        def frame(title):
            sl = new_slide()
            if os.path.exists(MARK):
                sl.shapes.add_picture(MARK, Inches(12.45), Inches(0.42), height=Inches(0.26))
            rule(sl, 0.7, 0.62, 0.85)
            t = tb(sl, 0.7, 0.85, 11.4, 1.0)
            rr = t.paragraphs[0].add_run(); rr.text = title
            rr.font.size, rr.font.bold, rr.font.color.rgb = Pt(27), True, TX
            return sl

        sl = frame(head)
        top = 1.95
        if lead:
            tl = tb(sl, 0.7, 1.88, 11.6, 0.95)
            rl = tl.paragraphs[0].add_run(); rl.text = lead
            rl.font.size, rl.font.color.rgb = Pt(13), TX2
            top = 2.85

        if groups:
            show = groups[:3]
            width = 11.9 / len(show)
            for i, (gh, items) in enumerate(show):
                g = tb(sl, 0.7 + i * width, top, width - 0.4, 4.1)
                first = True
                if gh:
                    pg = g.paragraphs[0]; rg = pg.add_run(); rg.text = gh.upper()
                    rg.font.size, rg.font.bold, rg.font.color.rgb = Pt(10.5), True, AC
                    first = False
                for it in items[:9]:
                    pi = g.paragraphs[0] if first else g.add_paragraph()
                    first = False; pi.space_before = Pt(7)
                    ri = pi.add_run(); ri.text = "— " + it
                    ri.font.size, ri.font.color.rgb = Pt(11.5), TX
            for gh, items in groups[3:]:
                s2 = frame(head + " (cont.)")
                g = tb(s2, 0.7, 1.95, 11.6, 4.6)
                pg = g.paragraphs[0]; rg = pg.add_run(); rg.text = gh.upper()
                rg.font.size, rg.font.bold, rg.font.color.rgb = Pt(10.5), True, AC
                for it in items:
                    pi = g.add_paragraph(); pi.space_before = Pt(8)
                    ri = pi.add_run(); ri.text = "— " + it
                    ri.font.size, ri.font.color.rgb = Pt(12), TX
        else:
            g = tb(sl, 0.7, top, 11.6, 4.2)
            first = True
            for t, x in blocks:
                if t == "p" and x != lead:
                    pi = g.paragraphs[0] if first else g.add_paragraph()
                    first = False; pi.space_before = Pt(10)
                    ri = pi.add_run(); ri.text = x
                    ri.font.size, ri.font.color.rgb = Pt(13), TX
    prs.save(PPTX)
    return sum(1 for _ in prs.slides)


# ------------------------------------------------------------------ DOCX
def page_bg(doc, hexcolor):
    bg = OxmlElement("w:background"); bg.set(qn("w:color"), hexcolor)
    doc.element.insert(0, bg)
    dp = OxmlElement("w:displayBackgroundShape")
    doc.settings.element.append(dp)


def hrule(doc, hexcolor="D97757"):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), hexcolor)
    bd.append(bot); pPr.append(bd)
    p.paragraph_format.space_after = DPt(2)
    return p


def build_docx(sections):
    d = docx.Document()
    page_bg(d, "FAF9F5")
    for s in d.sections:
        s.left_margin = s.right_margin = DIn(1.0)
        s.top_margin = DIn(0.85)
    f = d.styles["Normal"].font
    f.name, f.size, f.color.rgb = "Helvetica Neue", DPt(11), D_TX

    if os.path.exists(LOCKUP):
        d.add_picture(LOCKUP, height=DIn(0.30))
    hrule(d)

    sid, blocks = sections[0]
    p = d.add_paragraph(); r = p.add_run(heading_of(blocks))
    r.bold = True; r.font.size = DPt(21); r.font.color.rgb = D_TX
    p.paragraph_format.space_before = DPt(10)
    lead = next((x for t, x in blocks if t == "p"), "")
    if lead:
        pl = d.add_paragraph(); rl = pl.add_run(lead)
        rl.font.size = DPt(11.5); rl.font.color.rgb = D_TX2
    pe = d.add_paragraph(); re_ = pe.add_run("The Operating Review")
    re_.font.size = DPt(10.5); re_.bold = True; re_.font.color.rgb = D_AC

    for sid, blocks in sections[1:]:
        head = heading_of(blocks)
        lead = next((x for t, x in blocks if t == "p"), "")
        hp = d.add_paragraph(); hr = hp.add_run(head)
        hr.bold = True; hr.font.size = DPt(14.5); hr.font.color.rgb = D_TX
        hp.paragraph_format.space_before = DPt(18)
        hrule(d)
        if lead:
            pl = d.add_paragraph(); rl = pl.add_run(lead)
            rl.font.size = DPt(11); rl.font.color.rgb = D_TX2
        for gh, items in grouped(blocks, lead):
            if gh:
                pg = d.add_paragraph(); rg = pg.add_run(gh.upper())
                rg.bold = True; rg.font.size = DPt(10); rg.font.color.rgb = D_AC
                pg.paragraph_format.space_before = DPt(11)
            for it in items:
                pi = d.add_paragraph(it, style="List Bullet")
                for rr in pi.runs:
                    rr.font.size = DPt(11); rr.font.color.rgb = D_TX
    d.save(DOCX)
    return len(d.paragraphs)


if __name__ == "__main__":
    secs = parse(SRC)
    print(f"  parsed {len(secs)} sections")
    print(f"  pptx: {build_pptx(secs)} slides -> {PPTX}")
    print(f"  docx: {build_docx(secs)} paragraphs -> {DOCX}")
